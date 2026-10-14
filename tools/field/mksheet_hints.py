import json
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
D=json.load(open('/tmp/claude-0/sp/hints_now.json'))
AR,N,PRE,PATH,HI=D['AREAS'],D['HNOTE'],D['HPRE'],D['HPATH'],D['HINTS']
def area(f):
    for r in AR:
        if r[0]<=f<=r[1]: return r[2]
    return '(not named yet)'
def fmt(rooms):
    return ', '.join(str(f)+(' ('+'; '.join('%d, %d'%(q[0],q[1]) for q in sp)+')' if sp else '') for f,sp in rooms)
def rooms_for(k):   # same list the app's hint pop-up uses (showHint)
    if k in PATH: L=[(q[0],q[1]) for q in PATH[k]]
    else:
        pre=PRE.get(k,[]);rows=HI.get(k,[])
        L=[(q[0],q[1]) for q in pre]+[(r[0],[[r[2],r[3]]] if len(r)>3 else []) for r in rows[:max(1,3-len(pre))]]
    out=[]   # the same room twice in a row (two routines set the step) is listed once
    for f,sp in L:
        if out and out[-1][0]==f: out[-1]=(f,out[-1][1]+[q for q in sp if q not in out[-1][1]])
        else: out.append((f,list(sp)))
    return out
old=load_workbook('/root/.claude/uploads/8e6fc456-c737-554d-97c2-7a103c217ffc/93cea726-Xenogears_story_hints_1.74.xlsx').active   # Jack's edited copy
OLD={r[0]:r for r in old.iter_rows(min_row=4,values_only=True) if r[0] is not None}
PROP={p['step']:p for g in 'ABCD' for p in json.load(open('/tmp/claude-0/sp/prop_%s.json'%g))}
steps=sorted(set(map(int,HI))|set(map(int,N)))
wb=Workbook();ws=wb.active;ws.title='Story hints'
F=lambda **k:Font(name='Arial',**k)
ws['A1']='Xenogears Link 1.85 - story hints';ws['A1'].font=F(bold=True,size=14)
ws['A2']=('Edit the yellow columns. Rooms: list them in the order to visit, the last one is where the step happens; add a green spot as room (x, z) '
 'using the first and third "pos" numbers from DEV OPTIONS, e.g. 102 (1398, -44). Hint text: each sentence (ending in a full stop) is its own page in the app, '
 'paired with the rooms in order. Grey text = not yet checked in game. Leave Step as-is. Column I: hints rewritten in 1.85 from the GameFAQs walkthroughs (with how sure, and which guides), or where the app differs from your edited 1.74 sheet. Steps you edited were not changed.')
ws['A2'].font=F(italic=True,size=9);ws['A2'].alignment=Alignment(wrap_text=True,vertical='top');ws.merge_cells('A2:I2');ws.row_dimensions[2].height=40
H=['Step','Chapter (in game)','Area','Trigger','Rooms in order, with spots (edit)','Hint text (edit)','Checked (Y/N)','Your notes','What changed']
W=[7,22,20,20,34,60,10,40,30]
yel=PatternFill('solid',fgColor='FFF2CC');hd=PatternFill('solid',fgColor='1F3864');thin=Side(style='thin',color='BFBFBF')
for i,(h,w) in enumerate(zip(H,W),1):
    c=ws.cell(3,i,h);c.font=F(bold=True,color='FFFFFF');c.fill=hd;c.alignment=Alignment(wrap_text=True,vertical='center')
    ws.column_dimensions[c.column_letter].width=w
dv=DataValidation(type='list',formula1='"Y,N"',allow_blank=True);ws.add_data_validation(dv)
r=4;nch=0
for s in steps:
    k=str(s);note=N.get(k,'');guess=note.startswith('?') or not note;txt=note.lstrip('?')
    rl=rooms_for(k);o=OLD.get(s)
    ch=[]
    if o is None: ch.append('new step')
    else:
        if (o[4] or '')!=fmt(rl): ch.append('rooms')
        if (o[5] or '').strip()!=txt.strip(): ch.append('text')
        if (o[6]=='Y')!=(not guess): ch.append('checked')
    if s in PROP:
        q=PROP[s];ch=['1.85 from guides, '+q['confidence']+' ('+', '.join(q['sources'][:4])+')'+(', rooms' if q.get('rooms') else '')]
    nch+=bool(ch)
    vals=[s,o[1] if o else '',', '.join(dict.fromkeys(area(f) for f,_ in rl)) if rl else (o[2] if o else ''),o[3] if o else '',fmt(rl),txt,'N' if guess else 'Y',o[7] if o else None,', '.join(ch)]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v);c.font=F(size=10,color='808080' if guess and i==6 else '000000',bold=(i==9))
        c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(bottom=thin)
        if i in(5,6,7,8):c.fill=yel
    dv.add(ws.cell(r,7));r+=1
ws.freeze_panes='A4';ws.auto_filter.ref='A3:I%d'%(r-1)
out='/mnt/user-data/outputs/Xenogears story hints 1.85.xlsx'
import os;os.makedirs('/mnt/user-data/outputs',exist_ok=True);wb.save(out);print(r-4,'rows',nch,'changed',len(OLD),'old')
